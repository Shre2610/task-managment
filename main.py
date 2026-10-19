import sqlite3
import os
import datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import Flask, render_template, request, redirect, url_for, flash, send_file

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'supersecretkeyforflashingmessages')

DATABASE = os.environ.get('DATABASE_PATH', 'database.db')

VALID_PRIORITIES = ['Low', 'Medium', 'High', 'Urgent']
VALID_STATUSES = ['To Do', 'In Progress', 'Completed']
DEFAULT_CATEGORIES = ['Development', 'Design', 'Research', 'Testing', 'Documentation']

def get_db_connection():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn

def get_employee_id(name):
    mapping = {
        "Raj Sharma": "EMP1024",
        "Priya Mehta": "EMP1031",
        "Rohan Verma": "EMP1018",
        "Neha Kapoor": "EMP1042",
        "Karan Singh": "EMP1009",
        "Anjali Patel": "EMP1005"
    }
    return mapping.get(name, "EMP1099")

def get_employee_email(name):
    return f"{name.lower().replace(' ', '.')}@example.com" if name else ""

def format_date_display(date_str):
    if not date_str:
        return ""
    try:
        dt = datetime.datetime.strptime(date_str.split()[0], '%Y-%m-%d')
        return dt.strftime('%d %b %Y')
    except Exception:
        return date_str

def format_task_row(row):
    d = dict(row)
    assignee = d.get('assignee', '') or d.get('employee_name', '')
    d['assignee'] = assignee
    if not d.get('assignee_id'):
        d['assignee_id'] = get_employee_id(assignee)
    d['assignee_email'] = get_employee_email(assignee)

    # Format created_at
    created_at = d.get('created_at', '')
    if created_at:
        try:
            dt = datetime.datetime.strptime(created_at.split('.')[0], '%Y-%m-%d %H:%M:%S')
            d['formatted_created_at'] = dt.strftime('%d %b %Y, %I:%M %p')
        except Exception:
            d['formatted_created_at'] = created_at
    else:
        d['formatted_created_at'] = ""

    # Format dates
    due_date = d.get('due_date', '')
    d['formatted_due_date'] = format_date_display(due_date)
    d['formatted_start_date'] = format_date_display(d.get('start_date', ''))

    # Check overdue status
    today_str = datetime.date.today().isoformat()
    status = d.get('status', 'To Do')
    d['is_overdue'] = bool(due_date and due_date < today_str and status != 'Completed')
    d['is_due_today'] = bool(due_date and due_date == today_str and status != 'Completed')

    # Days calculation
    if due_date:
        try:
            due_dt = datetime.datetime.strptime(due_date, '%Y-%m-%d').date()
            today_dt = datetime.date.today()
            diff = (due_dt - today_dt).days
            d['days_diff'] = diff
            if d['is_overdue']:
                d['deadline_label'] = f"{abs(diff)}d overdue"
            elif diff == 0:
                d['deadline_label'] = "Due today"
            elif diff == 1:
                d['deadline_label'] = "Due tomorrow"
            else:
                d['deadline_label'] = f"In {diff} days"
        except Exception:
            d['days_diff'] = None
            d['deadline_label'] = ""
    else:
        d['days_diff'] = None
        d['deadline_label'] = ""

    return d

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()

    # 1. Preserve legacy allocations table structure if present
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS allocations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_name TEXT NOT NULL,
            project_name TEXT NOT NULL,
            project_type TEXT NOT NULL,
            month TEXT NOT NULL,
            allocation_percentage INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            employee_id TEXT
        )
    ''')

    # 2. Create tasks table with constraints and indexes
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            assignee TEXT NOT NULL,
            assignee_id TEXT,
            category TEXT NOT NULL DEFAULT 'Development',
            priority TEXT NOT NULL DEFAULT 'Medium' CHECK(priority IN ('Low', 'Medium', 'High', 'Urgent')),
            status TEXT NOT NULL DEFAULT 'To Do' CHECK(status IN ('To Do', 'In Progress', 'Completed')),
            start_date TEXT,
            due_date TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    conn.commit()

    # Create indexes for high-frequency queries
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_tasks_status ON tasks(status)")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_tasks_priority ON tasks(priority)")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_tasks_due_date ON tasks(due_date)")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_tasks_assignee ON tasks(assignee)")
    conn.commit()

    # 3. Seed initial task records if table is empty
    cursor.execute('SELECT COUNT(*) FROM tasks')
    count = cursor.fetchone()[0]
    if count == 0:
        today = datetime.date.today()
        # Seed realistic tasks with overdue, in-progress, and completed states
        sample_tasks = [
            ("Design Responsive Dashboard Wireframes", "Create low-fidelity and high-fidelity mockups for desktop and mobile views.", "Priya Mehta", "EMP1031", "Design", "High", "Completed", (today - datetime.timedelta(days=18)).isoformat(), (today - datetime.timedelta(days=7)).isoformat()),
            ("Implement User Authentication & Sessions", "Add secure password hashing, session tokens, and route protection middleware.", "Raj Sharma", "EMP1024", "Development", "Urgent", "In Progress", (today - datetime.timedelta(days=14)).isoformat(), (today - datetime.timedelta(days=3)).isoformat()),
            ("Set Up SQLite Schema & Constraints", "Define relational schema, check constraints, foreign keys, and indexes.", "Rohan Verma", "EMP1018", "Development", "High", "Completed", (today - datetime.timedelta(days=17)).isoformat(), (today - datetime.timedelta(days=9)).isoformat()),
            ("Draft API Documentation & Swagger Spec", "Document all REST endpoints, request payloads, and status codes.", "Neha Kapoor", "EMP1042", "Documentation", "Medium", "To Do", (today - datetime.timedelta(days=7)).isoformat(), (today + datetime.timedelta(days=6)).isoformat()),
            ("Conduct End-to-End Regression Testing", "Run comprehensive automated integration tests across all core workflows.", "Karan Singh", "EMP1009", "Testing", "Urgent", "In Progress", (today - datetime.timedelta(days=4)).isoformat(), (today + datetime.timedelta(days=3)).isoformat()),
            ("Analyze User Experience & Accessibility", "Audit contrast ratios, keyboard accessibility, and screen reader labels.", "Anjali Patel", "EMP1005", "Research", "Low", "To Do", (today - datetime.timedelta(days=1)).isoformat(), (today + datetime.timedelta(days=11)).isoformat()),
            ("Optimize Database Indexing & Query Latency", "Profile slow queries and verify coverage for compound filter criteria.", "Raj Sharma", "EMP1024", "Development", "High", "In Progress", (today - datetime.timedelta(days=5)).isoformat(), (today + datetime.timedelta(days=1)).isoformat()),
            ("Configure CI/CD Deployment Pipeline", "Set up automated test runs and build checks on pull request creation.", "Rohan Verma", "EMP1018", "Development", "Medium", "Completed", (today - datetime.timedelta(days=24)).isoformat(), (today - datetime.timedelta(days=14)).isoformat()),
            ("Implement Export to Excel for Reports", "Format downloadable spreadsheets with custom headers, borders, and filters.", "Priya Mehta", "EMP1031", "Development", "Medium", "Completed", (today - datetime.timedelta(days=11)).isoformat(), (today - datetime.timedelta(days=1)).isoformat()),
            ("Fix Mobile Navigation Drawer Overlay", "Resolve touch event bubbling issue on mobile sidebar backdrop.", "Karan Singh", "EMP1009", "Design", "High", "In Progress", (today - datetime.timedelta(days=3)).isoformat(), (today + datetime.timedelta(days=5)).isoformat()),
            ("Security Audit on Input Sanitization", "Verify parameterized queries and sanitize user-submitted inputs against XSS.", "Neha Kapoor", "EMP1042", "Testing", "Urgent", "To Do", today.isoformat(), (today + datetime.timedelta(days=7)).isoformat()),
            ("Write User Onboarding Walkthrough Guide", "Prepare quickstart manual and step-by-step screenshots for new users.", "Anjali Patel", "EMP1005", "Documentation", "Low", "To Do", today.isoformat(), (today + datetime.timedelta(days=17)).isoformat())
        ]
        cursor.executemany('''
            INSERT INTO tasks (title, description, assignee, assignee_id, category, priority, status, start_date, due_date)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', sample_tasks)
        conn.commit()

    conn.close()

def get_avatar_class(name):
    code = sum(ord(c) for c in name) % 5 if name else 0
    classes = ['avatar-teal', 'avatar-purple', 'avatar-orange', 'avatar-blue', 'avatar-yellow']
    return classes[code]

def get_initials(name):
    if not name:
        return "TF"
    parts = name.split()
    if len(parts) >= 2:
        return (parts[0][0] + parts[1][0]).upper()
    return name[:2].upper()

def get_priority_badge_class(priority):
    p = (priority or '').lower()
    if p == 'urgent':
        return 'badge-priority urgent'
    elif p == 'high':
        return 'badge-priority high'
    elif p == 'medium':
        return 'badge-priority medium'
    return 'badge-priority low'

def get_status_badge_class(status):
    s = (status or '').lower()
    if s == 'completed':
        return 'badge-status completed'
    elif s in ('in progress', 'inprogress'):
        return 'badge-status inprogress'
    return 'badge-status todo'

def get_category_badge_class(category):
    c = (category or '').lower()
    if 'dev' in c:
        return 'badge-cat dev'
    elif 'design' in c:
        return 'badge-cat design'
    elif 'research' in c or 'r&d' in c:
        return 'badge-cat research'
    elif 'test' in c:
        return 'badge-cat testing'
    elif 'doc' in c:
        return 'badge-cat docs'
    return 'badge-cat default'

@app.context_processor
def utility_processor():
    return {
        'get_avatar_class': get_avatar_class,
        'get_initials': get_initials,
        'get_priority_badge_class': get_priority_badge_class,
        'get_status_badge_class': get_status_badge_class,
        'get_category_badge_class': get_category_badge_class,
        'today_date': datetime.date.today().isoformat()
    }

# ==========================================
# ROUTES
# ==========================================

@app.route('/')
@app.route('/dashboard')
def dashboard():
    conn = get_db_connection()
    today_str = datetime.date.today().isoformat()

    total_tasks = conn.execute("SELECT COUNT(*) FROM tasks").fetchone()[0]
    todo_tasks = conn.execute("SELECT COUNT(*) FROM tasks WHERE status = 'To Do'").fetchone()[0]
    in_progress_tasks = conn.execute("SELECT COUNT(*) FROM tasks WHERE status = 'In Progress'").fetchone()[0]
    completed_tasks = conn.execute("SELECT COUNT(*) FROM tasks WHERE status = 'Completed'").fetchone()[0]
    overdue_tasks = conn.execute("SELECT COUNT(*) FROM tasks WHERE status != 'Completed' AND due_date < ?", (today_str,)).fetchone()[0]

    completion_rate = round((completed_tasks / total_tasks * 100)) if total_tasks > 0 else 0

    # Fetch recent 5 tasks
    recent_rows = conn.execute("SELECT * FROM tasks ORDER BY id DESC LIMIT 5").fetchall()
    recent_tasks = [format_task_row(r) for r in recent_rows]

    # Fetch urgent/overdue action items (up to 3)
    urgent_rows = conn.execute('''
        SELECT * FROM tasks
        WHERE status != 'Completed' AND (due_date < ? OR priority = 'Urgent')
        ORDER BY due_date ASC LIMIT 3
    ''', (today_str,)).fetchall()
    urgent_tasks = [format_task_row(r) for r in urgent_rows]

    # Category and Assignee choices for quick filters
    categories = [r['category'] for r in conn.execute('SELECT DISTINCT category FROM tasks ORDER BY category').fetchall()]
    assignees = [r['assignee'] for r in conn.execute('SELECT DISTINCT assignee FROM tasks ORDER BY assignee').fetchall()]

    conn.close()

    return render_template('dashboard.html',
                           total_tasks=total_tasks,
                           todo_tasks=todo_tasks,
                           in_progress_tasks=in_progress_tasks,
                           completed_tasks=completed_tasks,
                           overdue_tasks=overdue_tasks,
                           completion_rate=completion_rate,
                           recent_tasks=recent_tasks,
                           urgent_tasks=urgent_tasks,
                           categories=categories,
                           assignees=assignees)

@app.route('/tasks')
@app.route('/allocations') # Route alias for backward compatibility
def tasks():
    q = request.args.get('q', '').strip()
    status_filter = request.args.get('status', '').strip()
    priority_filter = request.args.get('priority', '').strip()
    assignee_filter = request.args.get('assignee', '').strip()
    # Support backward-compatible 'employee' parameter
    if not assignee_filter and request.args.get('employee'):
        assignee_filter = request.args.get('employee', '').strip()
    category_filter = request.args.get('category', '').strip()
    due_filter = request.args.get('due_filter', '').strip()
    sort_filter = request.args.get('sort', 'newest').strip()
    page = request.args.get('page', 1, type=int)

    query_str = "SELECT * FROM tasks WHERE 1=1"
    params = []

    today_str = datetime.date.today().isoformat()

    if q:
        query_str += " AND (title LIKE ? OR description LIKE ? OR assignee LIKE ? OR category LIKE ?)"
        params.extend([f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%"])

    if status_filter:
        query_str += " AND status = ?"
        params.append(status_filter)

    if priority_filter:
        query_str += " AND priority = ?"
        params.append(priority_filter)

    if assignee_filter:
        query_str += " AND assignee = ?"
        params.append(assignee_filter)

    if category_filter:
        query_str += " AND category = ?"
        params.append(category_filter)

    if due_filter == 'overdue':
        query_str += " AND due_date < ? AND status != 'Completed'"
        params.append(today_str)
    elif due_filter == 'today':
        query_str += " AND due_date = ?"
        params.append(today_str)
    elif due_filter == 'week':
        week_end = (datetime.date.today() + datetime.timedelta(days=7)).isoformat()
        query_str += " AND due_date >= ? AND due_date <= ?"
        params.extend([today_str, week_end])
    elif due_filter == 'month':
        month_end = (datetime.date.today() + datetime.timedelta(days=30)).isoformat()
        query_str += " AND due_date >= ? AND due_date <= ?"
        params.extend([today_str, month_end])

    # Sorting
    if sort_filter == 'due_asc':
        query_str += " ORDER BY due_date ASC"
    elif sort_filter == 'due_desc':
        query_str += " ORDER BY due_date DESC"
    elif sort_filter == 'priority':
        query_str += """ ORDER BY 
            CASE priority 
                WHEN 'Urgent' THEN 1 
                WHEN 'High' THEN 2 
                WHEN 'Medium' THEN 3 
                WHEN 'Low' THEN 4 
                ELSE 5 
            END ASC, due_date ASC"""
    elif sort_filter == 'title':
        query_str += " ORDER BY title ASC"
    else: # newest
        query_str += " ORDER BY id DESC"

    # Pagination
    limit = 6
    offset = max(0, (page - 1) * limit)

    conn = get_db_connection()
    count_query = query_str.replace("SELECT *", "SELECT COUNT(*)", 1)
    total_records = conn.execute(count_query, params).fetchone()[0]

    # Fetch paginated rows
    query_str += " LIMIT ? OFFSET ?"
    params.extend([limit, offset])
    rows = conn.execute(query_str, params).fetchall()
    tasks_list = [format_task_row(r) for r in rows]

    # Distinct values for dropdown filters
    all_assignees = [r['assignee'] for r in conn.execute('SELECT DISTINCT assignee FROM tasks ORDER BY assignee').fetchall()]
    all_categories = [r['category'] for r in conn.execute('SELECT DISTINCT category FROM tasks ORDER BY category').fetchall()]
    conn.close()

    total_pages = max(1, (total_records + limit - 1) // limit)
    showing_start = offset + 1 if total_records > 0 else 0
    showing_end = min(offset + limit, total_records)

    return render_template('allocations.html',
                           tasks_list=tasks_list,
                           all_assignees=all_assignees,
                           all_categories=all_categories,
                           valid_priorities=VALID_PRIORITIES,
                           valid_statuses=VALID_STATUSES,
                           q=q,
                           status_filter=status_filter,
                           priority_filter=priority_filter,
                           assignee_filter=assignee_filter,
                           category_filter=category_filter,
                           due_filter=due_filter,
                           sort_filter=sort_filter,
                           page=page,
                           total_pages=total_pages,
                           total_records=total_records,
                           showing_start=showing_start,
                           showing_end=showing_end)

@app.route('/add_task', methods=['GET', 'POST'])
@app.route('/add_assignment', methods=['GET', 'POST']) # Route alias for backward compatibility
def add_task():
    conn = get_db_connection()

    if request.method == 'POST':
        edit_id = request.form.get('edit_id', '').strip()
        title = request.form.get('title', '').strip()
        description = request.form.get('description', '').strip()
        assignee = request.form.get('assignee', '').strip() or request.form.get('employee_name', '').strip()
        category = request.form.get('category', '').strip() or request.form.get('project_type', '').strip() or 'Development'
        priority = request.form.get('priority', '').strip() or 'Medium'
        status = request.form.get('status', '').strip() or 'To Do'
        start_date = request.form.get('start_date', '').strip()
        due_date = request.form.get('due_date', '').strip()

        errors = []
        if not title:
            errors.append("Task title is required.")
        elif len(title) > 200:
            errors.append("Task title must be 200 characters or less.")

        if not assignee:
            errors.append("Assignee name is required.")

        if not due_date:
            errors.append("Due date is required.")
        else:
            try:
                datetime.datetime.strptime(due_date, '%Y-%m-%d')
            except ValueError:
                errors.append("Due date must be in a valid YYYY-MM-DD format.")

        if start_date:
            try:
                dt_start = datetime.datetime.strptime(start_date, '%Y-%m-%d')
                if due_date:
                    dt_due = datetime.datetime.strptime(due_date, '%Y-%m-%d')
                    if dt_start > dt_due:
                        errors.append("Start date cannot be after the due date.")
            except ValueError:
                errors.append("Start date must be in a valid YYYY-MM-DD format.")

        if priority not in VALID_PRIORITIES:
            errors.append(f"Invalid priority. Choose from: {', '.join(VALID_PRIORITIES)}.")

        if status not in VALID_STATUSES:
            errors.append(f"Invalid status. Choose from: {', '.join(VALID_STATUSES)}.")

        if errors:
            for error in errors:
                flash(error, 'error')
            record = {
                'id': edit_id,
                'title': title,
                'description': description,
                'assignee': assignee,
                'category': category,
                'priority': priority,
                'status': status,
                'start_date': start_date,
                'due_date': due_date
            }
            all_assignees = [r['assignee'] for r in conn.execute('SELECT DISTINCT assignee FROM tasks ORDER BY assignee').fetchall()]
            all_categories = [r['category'] for r in conn.execute('SELECT DISTINCT category FROM tasks ORDER BY category').fetchall()]
            conn.close()
            return render_template('add_assignment.html',
                                   record=record,
                                   all_assignees=all_assignees,
                                   all_categories=all_categories or DEFAULT_CATEGORIES,
                                   valid_priorities=VALID_PRIORITIES,
                                   valid_statuses=VALID_STATUSES)

        assignee_id = get_employee_id(assignee)
        now_ts = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        if edit_id and edit_id != '-1':
            conn.execute('''
                UPDATE tasks
                SET title = ?, description = ?, assignee = ?, assignee_id = ?,
                    category = ?, priority = ?, status = ?, start_date = ?, due_date = ?, updated_at = ?
                WHERE id = ?
            ''', (title, description, assignee, assignee_id, category, priority, status, start_date or None, due_date, now_ts, edit_id))
            conn.commit()
            flash(f"Task #{edit_id} \"{title}\" updated successfully.", "success")
        else:
            cursor = conn.cursor()
            cursor.execute('''
                INSERT INTO tasks (title, description, assignee, assignee_id, category, priority, status, start_date, due_date, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (title, description, assignee, assignee_id, category, priority, status, start_date or None, due_date, now_ts, now_ts))
            new_id = cursor.lastrowid
            conn.commit()
            flash(f"New task #{new_id} \"{title}\" created successfully.", "success")

        conn.close()
        return redirect(url_for('tasks'))

    # GET request
    edit_id = request.args.get('edit')
    record = None
    if edit_id:
        row = conn.execute('SELECT * FROM tasks WHERE id = ?', (edit_id,)).fetchone()
        if row:
            record = dict(row)

    all_assignees = [r['assignee'] for r in conn.execute('SELECT DISTINCT assignee FROM tasks ORDER BY assignee').fetchall()]
    all_categories = [r['category'] for r in conn.execute('SELECT DISTINCT category FROM tasks ORDER BY category').fetchall()]
    conn.close()

    return render_template('add_assignment.html',
                           record=record,
                           all_assignees=all_assignees,
                           all_categories=all_categories or DEFAULT_CATEGORIES,
                           valid_priorities=VALID_PRIORITIES,
                           valid_statuses=VALID_STATUSES)

@app.route('/task/<int:id>/status', methods=['POST'])
def update_task_status(id):
    new_status = request.form.get('status', '').strip()
    if new_status not in VALID_STATUSES:
        flash(f"Invalid status: {new_status}", "error")
        return redirect(request.referrer or url_for('tasks'))

    conn = get_db_connection()
    row = conn.execute('SELECT title FROM tasks WHERE id = ?', (id,)).fetchone()
    if row:
        now_ts = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        conn.execute('UPDATE tasks SET status = ?, updated_at = ? WHERE id = ?', (new_status, now_ts, id))
        conn.commit()
        flash(f"Status for task #{id} changed to '{new_status}'.", "success")
    else:
        flash(f"Task #{id} not found.", "error")
    conn.close()
    return redirect(request.referrer or url_for('tasks'))

@app.route('/delete_task/<int:id>', methods=['POST'])
@app.route('/delete_assignment/<int:id>', methods=['POST']) # Route alias for backward compatibility
def delete_task(id):
    conn = get_db_connection()
    row = conn.execute('SELECT title FROM tasks WHERE id = ?', (id,)).fetchone()
    if row:
        title = row['title']
        conn.execute('DELETE FROM tasks WHERE id = ?', (id,))
        conn.commit()
        flash(f"Task #{id} \"{title}\" deleted successfully.", "success")
    else:
        flash(f"Task #{id} not found.", "error")
    conn.close()
    return redirect(request.referrer or url_for('tasks'))

@app.route('/assignee')
@app.route('/employee') # Route alias for backward compatibility
def assignee_view():
    name_param = request.args.get('name', '').strip()
    conn = get_db_connection()

    all_assignees = [r['assignee'] for r in conn.execute('SELECT DISTINCT assignee FROM tasks ORDER BY assignee').fetchall()]
    if not name_param and all_assignees:
        name_param = all_assignees[0]
    elif not name_param:
        name_param = "Raj Sharma"

    rows = conn.execute('SELECT * FROM tasks WHERE assignee LIKE ? ORDER BY due_date ASC', (f"%{name_param}%",)).fetchall()
    tasks_list = [format_task_row(r) for r in rows]

    matched_name = tasks_list[0]['assignee'] if tasks_list else name_param
    matched_id = tasks_list[0]['assignee_id'] if tasks_list else get_employee_id(name_param)
    matched_email = tasks_list[0]['assignee_email'] if tasks_list else get_employee_email(name_param)

    # Calculate assignee statistics
    today_str = datetime.date.today().isoformat()
    total_assigned = len(tasks_list)
    completed_count = sum(1 for t in tasks_list if t['status'] == 'Completed')
    in_progress_count = sum(1 for t in tasks_list if t['status'] == 'In Progress')
    overdue_count = sum(1 for t in tasks_list if t['is_overdue'])

    recent_rows = conn.execute('SELECT * FROM tasks ORDER BY id DESC LIMIT 5').fetchall()
    all_tasks_preview = [format_task_row(r) for r in recent_rows]
    conn.close()

    return render_template('employee.html',
                           name=matched_name,
                           emp_id=matched_id,
                           email=matched_email,
                           tasks_list=tasks_list,
                           total_assigned=total_assigned,
                           completed_count=completed_count,
                           in_progress_count=in_progress_count,
                           overdue_count=overdue_count,
                           all_tasks_preview=all_tasks_preview,
                           all_assignees=all_assignees)

@app.route('/timeline')
@app.route('/month') # Route alias for backward compatibility
def timeline_view():
    month_filter = request.args.get('month', '').strip()
    conn = get_db_connection()
    today = datetime.date.today()

    if not month_filter:
        month_filter = today.strftime('%Y-%m') # e.g. "2026-10"

    month_name = month_filter
    try:
        dt = datetime.datetime.strptime(month_filter, '%Y-%m')
        month_name = dt.strftime('%B %Y')
    except Exception:
        pass

    rows = conn.execute("SELECT * FROM tasks WHERE due_date LIKE ? ORDER BY due_date ASC", (f"{month_filter}%",)).fetchall()
    month_tasks = [format_task_row(r) for r in rows]

    today_str = today.isoformat()
    total_tasks = len(month_tasks)
    completed_count = sum(1 for t in month_tasks if t['status'] == 'Completed')
    in_progress_count = sum(1 for t in month_tasks if t['status'] == 'In Progress')
    overdue_count = sum(1 for t in month_tasks if t['is_overdue'])

    # Get all distinct months from task due dates
    dates = [r['due_date'] for r in conn.execute('SELECT DISTINCT due_date FROM tasks WHERE due_date IS NOT NULL ORDER BY due_date').fetchall()]
    all_months = sorted(list(set(d[:7] for d in dates if len(d) >= 7)))
    if not all_months:
        all_months = [today.strftime('%Y-%m')]

    conn.close()

    return render_template('month.html',
                           month_key=month_filter,
                           month_name=month_name,
                           total_tasks=total_tasks,
                           completed_count=completed_count,
                           in_progress_count=in_progress_count,
                           overdue_count=overdue_count,
                           month_tasks=month_tasks,
                           all_months=all_months)

@app.route('/search')
def search():
    name = request.args.get('name', '').strip()
    q = request.args.get('q', '').strip()
    search_term = q or name
    if not search_term:
        return redirect(url_for('dashboard'))

    # Check if search term matches a known assignee
    conn = get_db_connection()
    matched_assignee = conn.execute("SELECT assignee FROM tasks WHERE assignee LIKE ? LIMIT 1", (f"%{search_term}%",)).fetchone()
    conn.close()

    if matched_assignee:
        return redirect(url_for('assignee_view', name=matched_assignee['assignee']))
    return redirect(url_for('tasks', q=search_term))

@app.route('/export_excel')
def export_excel():
    q = request.args.get('q', '').strip()
    status_filter = request.args.get('status', '').strip()
    priority_filter = request.args.get('priority', '').strip()
    assignee_filter = request.args.get('assignee', '').strip() or request.args.get('employee', '').strip()
    category_filter = request.args.get('category', '').strip()
    name_filter = request.args.get('name', '').strip()
    month_filter = request.args.get('month', '').strip()

    query_str = "SELECT * FROM tasks WHERE 1=1"
    params = []

    if name_filter:
        query_str += " AND assignee LIKE ?"
        params.append(f"%{name_filter}%")

    if month_filter:
        query_str += " AND due_date LIKE ?"
        params.append(f"{month_filter}%")

    if q:
        query_str += " AND (title LIKE ? OR description LIKE ? OR assignee LIKE ? OR category LIKE ?)"
        params.extend([f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%"])

    if status_filter:
        query_str += " AND status = ?"
        params.append(status_filter)

    if priority_filter:
        query_str += " AND priority = ?"
        params.append(priority_filter)

    if assignee_filter:
        query_str += " AND assignee = ?"
        params.append(assignee_filter)

    if category_filter:
        query_str += " AND category = ?"
        params.append(category_filter)

    query_str += " ORDER BY due_date ASC, priority ASC"

    conn = get_db_connection()
    rows = conn.execute(query_str, params).fetchall()
    conn.close()

    if not rows:
        flash("No task records available to export with current filters.", "error")
        return redirect(request.referrer or url_for('tasks'))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tasks Report"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Task ID",
        "Title",
        "Category",
        "Assignee",
        "Priority",
        "Status",
        "Start Date",
        "Due Date",
        "Overdue?",
        "Created Date"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Teal 700
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row_idx, r in enumerate(rows, 2):
        formatted = format_task_row(r)
        overdue_text = "YES (OVERDUE)" if formatted.get('is_overdue') else "No"

        row_values = [
            formatted.get('id', ''),
            formatted.get('title', ''),
            formatted.get('category', ''),
            formatted.get('assignee', ''),
            formatted.get('priority', ''),
            formatted.get('status', ''),
            formatted.get('formatted_start_date', ''),
            formatted.get('formatted_due_date', ''),
            overdue_text,
            formatted.get('formatted_created_at', '')
        ]
        ws.append(row_values)

        for col_idx in range(1, 11):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border

            if col_idx in [1, 5, 6, 7, 8, 9, 10]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Highlight priority
            if col_idx == 5:
                if formatted.get('priority') == 'Urgent':
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B91C1C")
                elif formatted.get('priority') == 'High':
                    cell.font = Font(name="Calibri", size=10, bold=True, color="C2410C")
            # Highlight status
            elif col_idx == 6:
                if formatted.get('status') == 'Completed':
                    cell.font = Font(name="Calibri", size=10, bold=True, color="15803D")
                elif formatted.get('status') == 'In Progress':
                    cell.font = Font(name="Calibri", size=10, bold=True, color="1D4ED8")
            # Highlight overdue
            elif col_idx == 9 and formatted.get('is_overdue'):
                cell.font = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"taskflow_export_{timestamp}.xlsx"

    return send_file(
        file_stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

init_db()

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
